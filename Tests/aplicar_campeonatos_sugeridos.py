from openpyxl import load_workbook
from pathlib import Path


# =========================================================
# ARQUIVOS
# =========================================================

pasta = Path(__file__).resolve().parent

arquivo_base = pasta / "Banco_Jogadores_FINAL_ATUALIZADO.xlsx"
arquivo_sugestoes = pasta / "Sugestoes_Campeonatos.xlsx"

arquivo_saida = pasta / "Banco_Jogadores_COMPLETO.xlsx"


# =========================================================
# ABRE BASE PRINCIPAL
# =========================================================

wb_base = load_workbook(arquivo_base)
ws_base = wb_base.active


cab_base = {}

for celula in ws_base[1]:
    if celula.value:
        cab_base[str(celula.value).strip()] = celula.column


col_id_base = cab_base["ID_Jogador"]
col_campeonato_base = cab_base["Campeonato"]


# =========================================================
# ABRE PLANILHA DE SUGESTÕES
# =========================================================

wb_sug = load_workbook(arquivo_sugestoes)
ws_sug = wb_sug.active


cab_sug = {}

for celula in ws_sug[1]:
    if celula.value:
        cab_sug[str(celula.value).strip()] = celula.column


col_id_sug = cab_sug["ID_Jogador"]
col_sugestao = cab_sug["Campeonato_Sugerido"]


# =========================================================
# LÊ TODOS OS JOGADORES DA PLANILHA DE REVISÃO
# =========================================================

resultados = {}


for linha in range(2, ws_sug.max_row + 1):

    id_jogador = ws_sug.cell(
        linha,
        col_id_sug
    ).value

    campeonato = ws_sug.cell(
        linha,
        col_sugestao
    ).value


    if id_jogador is None:
        continue


    if campeonato:
        resultados[id_jogador] = campeonato
    else:
        resultados[id_jogador] = "NÃO ENCONTRADO"


# =========================================================
# APLICA NA BASE PRINCIPAL
# =========================================================

atualizados = 0
nao_encontrados = 0


for linha in range(2, ws_base.max_row + 1):

    id_jogador = ws_base.cell(
        linha,
        col_id_base
    ).value


    # Só mexe nos jogadores que estavam
    # na planilha Sugestoes_Campeonatos
    if id_jogador not in resultados:
        continue


    campeonato_novo = resultados[id_jogador]


    ws_base.cell(
        linha,
        col_campeonato_base
    ).value = campeonato_novo


    if campeonato_novo == "NÃO ENCONTRADO":

        nao_encontrados += 1

        print(
            f"ID {id_jogador}: NÃO ENCONTRADO"
        )

    else:

        atualizados += 1

        print(
            f"ID {id_jogador}: {campeonato_novo}"
        )


# =========================================================
# SALVA
# =========================================================

wb_base.save(
    arquivo_saida
)


print("\n================================")
print("PROCESSO FINALIZADO")
print("================================")

print(
    f"Campeonatos atualizados: {atualizados}"
)

print(
    f"Marcados como NÃO ENCONTRADO: {nao_encontrados}"
)

print("\nArquivo criado:")
print(arquivo_saida)