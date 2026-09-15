import re
import time
import random
import requests

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pathlib import Path
from urllib.parse import urljoin


# =========================================================
# CONFIGURAÇÕES
# =========================================================

arquivo_entrada = "Banco_Jogadores_Cristiano.xlsx"
arquivo_saida = "Banco_Jogadores_FINAL_ATUALIZADO.xlsx"

TEMPO_MINIMO = 2.0
TEMPO_MAXIMO = 4.0


# =========================================================
# CAMINHOS
# =========================================================

pasta = Path(__file__).resolve().parent

caminho_entrada = pasta / arquivo_entrada
caminho_saida = pasta / arquivo_saida


if not caminho_entrada.exists():

    print("Arquivo não encontrado:")
    print(caminho_entrada)

    raise SystemExit


# =========================================================
# ABRE EXCEL
# =========================================================

wb = load_workbook(caminho_entrada)
ws = wb.active


# =========================================================
# LOCALIZA COLUNAS
# =========================================================

cabecalhos = {}

for celula in ws[1]:

    if celula.value:

        cabecalhos[
            str(celula.value).strip()
        ] = celula.column


col_nome = cabecalhos.get("Nome")
col_clube = cabecalhos.get("Clube_Atual")
col_campeonato = cabecalhos.get("Campeonato")
col_link = cabecalhos.get("Link_Ogol")


if not all([
    col_nome,
    col_clube,
    col_campeonato,
    col_link
]):

    print("ERRO: alguma coluna não foi encontrada.")

    print(cabecalhos)

    raise SystemExit


# =========================================================
# CONEXÃO
# =========================================================

session = requests.Session()

session.headers.update({

    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/151.0.0.0 Safari/537.36"
    ),

    "Accept-Language": "pt-BR,pt;q=0.9"
})


# =========================================================
# LIMPA NOME DA COMPETIÇÃO
# =========================================================

def limpar_campeonato(nome):

    if not nome:
        return None

    nome = nome.strip()

    # Remove o ano final
    nome = re.sub(
        r"\s+20\d{2}\s*$",
        "",
        nome
    )

    # =====================================================
    # PADRONIZA ALGUNS NOMES
    # =====================================================

    nome_upper = nome.upper()


    if "SÉRIE A" in nome_upper:
        return "SÉRIE A"

    if "SÉRIE B" in nome_upper:
        return "SÉRIE B"

    if "SÉRIE C" in nome_upper:
        return "SÉRIE C"

    if "SÉRIE D" in nome_upper:
        return "SÉRIE D"


    if (
        "GAÚCHO" in nome_upper
        and (
            "DIVISÃO DE ACESSO" in nome_upper
            or "ACESSO" in nome_upper
        )
    ):
        return "GAÚCHO ACESSO"


    # Nos outros campeonatos,
    # apenas remove "Campeonato" do começo

    nome = re.sub(
        r"^Campeonato\s+",
        "",
        nome,
        flags=re.IGNORECASE
    )

    return nome.upper()


# =========================================================
# PEGA NOME NA PÁGINA DA EDIÇÃO
# =========================================================

def buscar_nome_edicao(url_edicao):

    try:

        resposta = session.get(
            url_edicao,
            timeout=20
        )

    except requests.RequestException:

        return None


    if resposta.status_code != 200:

        return None


    soup = BeautifulSoup(
        resposta.text,
        "html.parser"
    )


    # Primeiro tenta H1

    h1 = soup.find("h1")

    if h1:

        nome = h1.get_text(
            " ",
            strip=True
        )

        if nome:
            return limpar_campeonato(
                nome
            )


    # Se não tiver H1, usa título

    if soup.title:

        titulo = soup.title.get_text(
            " ",
            strip=True
        )

        # Ex:
        # Campeonato Gaúcho Divisão de Acesso 2026
        # - Classificações...

        titulo = titulo.split(
            " - "
        )[0]

        return limpar_campeonato(
            titulo
        )


    return None


# =========================================================
# DESCOBRE A EDIÇÃO ASSOCIADA AO CLUBE ATUAL
# =========================================================

def buscar_campeonato(
    url_jogador,
    clube_atual
):

    try:

        resposta = session.get(
            url_jogador,
            timeout=20
        )

    except requests.RequestException:

        return None, None


    if resposta.status_code != 200:

        print(
            f"   HTTP jogador: {resposta.status_code}"
        )

        return None, None


    soup = BeautifulSoup(
        resposta.text,
        "html.parser"
    )


    clube_lower = clube_atual.lower()


    # =====================================================
    # PROCURA LINKS DE EDIÇÃO QUE ESTEJAM
    # PRÓXIMOS AO CLUBE ATUAL
    # =====================================================

    candidatos = []


    for link in soup.find_all(
        "a",
        href=True
    ):

        href = link.get(
            "href",
            ""
        )


        if "edition.php?id_edicao=" not in href:
            continue


        # Sobe na estrutura HTML procurando
        # um bloco que também contenha o clube atual

        elemento = link

        encontrou_clube = False


        for _ in range(8):

            elemento = elemento.parent

            if elemento is None:
                break


            texto_bloco = elemento.get_text(
                " ",
                strip=True
            )


            if clube_lower in texto_bloco.lower():

                encontrou_clube = True
                break


        if encontrou_clube:

            url_edicao = urljoin(
                url_jogador,
                href
            )


            if url_edicao not in candidatos:

                candidatos.append(
                    url_edicao
                )


    # =====================================================
    # TESTA AS EDIÇÕES ENCONTRADAS
    # =====================================================

    for url_edicao in candidatos:

        campeonato = buscar_nome_edicao(
            url_edicao
        )


        if campeonato:

            return campeonato, url_edicao


    return None, None


# =========================================================
# PROCESSAMENTO
# =========================================================

total = 0
atualizados = 0
nao_encontrados = 0
sem_link = 0


for linha in range(
    2,
    ws.max_row + 1
):

    nome = ws.cell(
        linha,
        col_nome
    ).value


    if not nome:
        continue


    total += 1


    clube = ws.cell(
        linha,
        col_clube
    ).value


    link = ws.cell(
        linha,
        col_link
    ).value


    campeonato_antigo = ws.cell(
        linha,
        col_campeonato
    ).value


    print("\n================================")

    print(
        f"[{total}] {nome}"
    )

    print(
        f"   Clube atual: {clube}"
    )

    print(
        f"   Campeonato antigo: {campeonato_antigo}"
    )


    if not link:

        print(
            "   Sem Link_Ogol."
        )

        sem_link += 1
        continue


    if not clube:

        print(
            "   Sem Clube_Atual."
        )

        nao_encontrados += 1
        continue


    campeonato_novo, url_edicao = buscar_campeonato(
        str(link).strip(),
        str(clube).strip()
    )


    if campeonato_novo:

        ws.cell(
            linha,
            col_campeonato
        ).value = campeonato_novo


        print(
            f"   Campeonato atual: {campeonato_novo}"
        )

        print(
            f"   Edição: {url_edicao}"
        )


        atualizados += 1


    else:

        print(
            "   Campeonato não encontrado."
        )

        # NÃO APAGA o valor antigo
        nao_encontrados += 1


    # =====================================================
    # SALVA PROGRESSO
    # =====================================================

    try:

        wb.save(
            caminho_saida
        )

    except PermissionError:

        print(
            "\nArquivo bloqueado pelo Windows/OneDrive."
        )

        print(
            "Encerrando para não perder o progresso."
        )

        raise


    time.sleep(
        random.uniform(
            TEMPO_MINIMO,
            TEMPO_MAXIMO
        )
    )


# =========================================================
# FINAL
# =========================================================

wb.save(
    caminho_saida
)


print("\n================================")
print("PROCESSO FINALIZADO")
print("================================")

print(
    f"Jogadores analisados: {total}"
)

print(
    f"Campeonatos atualizados: {atualizados}"
)

print(
    f"Não encontrados: {nao_encontrados}"
)

print(
    f"Sem Link_Ogol: {sem_link}"
)

print("\nArquivo criado:")

print(
    caminho_saida
)