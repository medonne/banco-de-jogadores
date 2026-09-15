import time
import random
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pathlib import Path


# =========================================================
# CONFIGURAÇÕES
# =========================================================

arquivo_entrada = "Banco_Jogadores_CLUBES_ATUALIZADOS.xlsx"
arquivo_saida = "Banco_Jogadores_CLUBES_ATUALIZADOS.xlsx"

TEMPO_MINIMO = 2.0
TEMPO_MAXIMO = 4.0


# =========================================================
# ARQUIVOS
# =========================================================

pasta = Path(__file__).resolve().parent

caminho_entrada = pasta / arquivo_entrada
caminho_saida = pasta / arquivo_saida

if not caminho_entrada.exists():
    print("ERRO: arquivo não encontrado:")
    print(caminho_entrada)
    input("\nPressione ENTER para sair...")
    raise SystemExit


# =========================================================
# ABRE PLANILHA
# =========================================================

wb = load_workbook(caminho_entrada)
ws = wb.active


# =========================================================
# LOCALIZA COLUNAS PELO CABEÇALHO
# =========================================================

cabecalhos = {}

for celula in ws[1]:
    if celula.value:
        cabecalhos[str(celula.value).strip()] = celula.column


col_nome = cabecalhos.get("Nome")
col_clube = cabecalhos.get("Clube_Atual")
col_pais = cabecalhos.get("Pais")
col_link = cabecalhos.get("Link_Ogol")


if not col_nome:
    raise ValueError("Coluna 'Nome' não encontrada.")

if not col_clube:
    raise ValueError("Coluna 'Clube_Atual' não encontrada.")

if not col_pais:
    raise ValueError("Coluna 'Pais' não encontrada.")

if not col_link:
    raise ValueError("Coluna 'Link_Ogol' não encontrada.")


# =========================================================
# CONEXÃO
# =========================================================

session = requests.Session()

session.headers.update({
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/151.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "pt-BR,pt;q=0.9"
})


# =========================================================
# BUSCA CLUBE ATUAL + PAÍS DO CLUBE
# =========================================================

def buscar_clube_pais(url):

    try:
        resposta = session.get(url, timeout=20)

    except requests.RequestException as erro:
        print(f"   Erro de conexão: {erro}")
        return "ERRO_CONEXAO", None

    if resposta.status_code in (403, 429, 500, 502, 503, 504):
        print(f"   OGol indisponível/bloqueando. HTTP {resposta.status_code}")
        return "PARAR", None

    if resposta.status_code != 200:
        print(f"   HTTP {resposta.status_code}")
        return None, None


    soup = BeautifulSoup(
        resposta.text,
        "html.parser"
    )


    # =====================================================
    # PROCURA A LINHA "Clube atual"
    # =====================================================

    linhas_dados = soup.select(".card-data__row")

    for linha in linhas_dados:

        label = linha.select_one(".card-data__label")

        if not label:
            continue

        nome_label = label.get_text(
            " ",
            strip=True
        ).lower()

        if nome_label != "clube atual":
            continue


        # =================================================
        # CLUBE
        # =================================================

        clube = None

        texto_clube = linha.select_one(
            ".micrologo_and_text .text"
        )

        if texto_clube:

            clube = texto_clube.get_text(
                " ",
                strip=True
            )


        # =================================================
        # PAÍS DO CLUBE
        #
        # Exemplo:
        # <a href="/pais/brasil" title="Brasil">
        # =================================================

        pais = None

        link_pais = linha.select_one(
            'a[href*="/pais/"]'
        )

        if link_pais:

            pais = link_pais.get(
                "title"
            )

            if not pais:

                href = link_pais.get(
                    "href",
                    ""
                )

                if "/pais/" in href:

                    pais = href.split(
                        "/pais/"
                    )[-1]

                    pais = pais.replace(
                        "-",
                        " "
                    ).title()


        return clube, pais


    return None, None


# =========================================================
# PROCESSAMENTO
# =========================================================

total = 0
atualizados = 0
nao_encontrados = 0
sem_link = 0


for linha in range(
    2,
    ws.max_row + 1
):

    nome = ws.cell(
        linha,
        col_nome
    ).value

    link = ws.cell(
        linha,
        col_link
    ).value

    if not nome:
        continue

    total += 1

# Já processamos os 305 primeiros jogadores
if total < 306:
    continue


if not link:

        print(
            f"[{total}] {nome}: SEM LINK OGOL"
        )

        sem_link += 1
        continue


    clube_antigo = ws.cell(
        linha,
        col_clube
    ).value

    pais_antigo = ws.cell(
        linha,
        col_pais
    ).value


    print("\n--------------------------------")

    print(
        f"[{total}] {nome}"
    )

    print(
        f"   Clube antigo: {clube_antigo}"
    )

    print(
        f"   País antigo: {pais_antigo}"
    )


    clube_novo, pais_novo = buscar_clube_pais(
        str(link).strip()
    )


    # =====================================================
    # BLOQUEIO DO SITE
    # =====================================================

    if clube_novo == "PARAR":

        print(
            "\nOGol ficou indisponível."
        )

        print(
            "Salvando progresso..."
        )

        wb.save(
            caminho_saida
        )

        break


    # =====================================================
    # ERRO DE CONEXÃO
    # =====================================================

    if clube_novo == "ERRO_CONEXAO":

        nao_encontrados += 1

        continue


    # =====================================================
    # ATUALIZA
    # =====================================================

    if clube_novo:

        clube_novo = clube_novo.strip()

        ws.cell(
            linha,
            col_clube
        ).value = clube_novo

        print(
            f"   Clube atual OGol: {clube_novo}"
        )


        if pais_novo:

            pais_novo = pais_novo.strip().upper()

            ws.cell(
                linha,
                col_pais
            ).value = pais_novo

            print(
                f"   País do clube: {pais_novo}"
            )

        else:

            print(
                "   País do clube não encontrado."
            )


        atualizados += 1

        # Salva a cada jogador
        wb.save(
            caminho_saida
        )


    else:

        print(
            "   Clube atual não encontrado."
        )

        nao_encontrados += 1


    # =====================================================
    # INTERVALO
    # =====================================================

    time.sleep(
        random.uniform(
            TEMPO_MINIMO,
            TEMPO_MAXIMO
        )
    )


else:

    wb.save(
        caminho_saida
    )


# =========================================================
# RESUMO
# =========================================================

print("\n================================")
print("PROCESSO FINALIZADO")
print("================================")

print(
    f"Jogadores analisados: {total}"
)

print(
    f"Atualizados: {atualizados}"
)

print(
    f"Não encontrados: {nao_encontrados}"
)

print(
    f"Sem Link_Ogol: {sem_link}"
)

print("\nArquivo criado:")

print(
    caminho_saida
)