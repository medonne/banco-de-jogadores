import requests
import re
import time
import shutil

from bs4 import BeautifulSoup
from datetime import datetime
from openpyxl import load_workbook


# ==========================================================
# CONFIGURAÇÕES
# ==========================================================

arquivo_entrada = "Banco_Jogadores_ATUALIZADO.xlsx"
arquivo_saida = "Banco_Jogadores_FINAL.xlsx"

tentativas_maximas = 3
pausa_entre_jogadores = 3

headers = {
    "User-Agent": "Mozilla/5.0"
}


# ==========================================================
# TESTAR SE O OGOL ESTÁ DISPONÍVEL
# ==========================================================

def ogol_disponivel():
    try:
        resposta = requests.get(
            "https://www.ogol.com.br/",
            headers=headers,
            timeout=15
        )

        if resposta.status_code != 200:
            return False

        texto = resposta.text.lower()

        if "serviço temporariamente suspenso" in texto:
            return False

        if "servico temporariamente suspenso" in texto:
            return False

        return True

    except requests.exceptions.RequestException:
        return False


print()
print("=" * 60)
print("VERIFICANDO DISPONIBILIDADE DO OGOL")
print("=" * 60)

if not ogol_disponivel():
    print()
    print("OGol está indisponível no momento.")
    print("Nenhuma coleta será realizada.")
    print("Tente novamente mais tarde.")
    print()
    raise SystemExit

print()
print("OGol disponível ✓")
print()


# ==========================================================
# PREPARAR ARQUIVO
# ==========================================================

# Se o arquivo final ainda não existe, cria uma cópia
try:
    wb = load_workbook(arquivo_saida)
    print("Retomando arquivo existente:", arquivo_saida)

except FileNotFoundError:
    shutil.copy2(
        arquivo_entrada,
        arquivo_saida
    )

    wb = load_workbook(arquivo_saida)

    print("Novo arquivo criado:", arquivo_saida)


ws = wb.active


# ==========================================================
# IDENTIFICAR COLUNAS
# ==========================================================

cabecalhos = {}

for celula in ws[1]:
    if celula.value is not None:
        cabecalhos[str(celula.value).strip()] = celula.column


col_nome = cabecalhos["Nome"]
col_data = cabecalhos["Data_Nascimento"]
col_link = cabecalhos["Link_Ogol"]


# ==========================================================
# IDENTIFICAR JOGADORES PENDENTES
# ==========================================================

linhas_pendentes = []

for linha in range(2, ws.max_row + 1):

    nome = ws.cell(
        linha,
        col_nome
    ).value

    data = ws.cell(
        linha,
        col_data
    ).value

    if nome and not data:
        linhas_pendentes.append(linha)


total = len(linhas_pendentes)

print()
print("=" * 60)
print("JOGADORES PENDENTES")
print("=" * 60)

print(f"Total sem data: {total}")

if total == 0:
    print("Todos os jogadores já possuem data de nascimento.")
    wb.close()
    raise SystemExit


# ==========================================================
# SESSÃO HTTP
# ==========================================================

session = requests.Session()
session.headers.update(headers)


encontradas = 0
nao_encontradas = 0
erros = 0


# ==========================================================
# COLETA
# ==========================================================

for contador, linha in enumerate(
    linhas_pendentes,
    start=1
):

    nome = ws.cell(
        linha,
        col_nome
    ).value

    url = ws.cell(
        linha,
        col_link
    ).value


    if not url:

        print(
            f"[{contador}/{total}] "
            f"{nome} -> LINK AUSENTE"
        )

        nao_encontradas += 1
        continue


    sucesso = False


    for tentativa in range(
        1,
        tentativas_maximas + 1
    ):

        try:

            resposta = session.get(
                str(url),
                timeout=30
            )


            # ==================================================
            # SITE FORA DO AR
            # ==================================================

            if resposta.status_code == 503:

                print()
                print(
                    f"OGol retornou erro 503 ao acessar {nome}."
                )

                print(
                    "O site parece estar temporariamente "
                    "indisponível."
                )

                print(
                    "A coleta será interrompida com segurança."
                )

                wb.save(arquivo_saida)
                wb.close()

                raise SystemExit


            if resposta.status_code == 429:

                print(
                    f"[{contador}/{total}] "
                    f"{nome} -> muitas requisições"
                )

                print(
                    "Aguardando 60 segundos..."
                )

                time.sleep(60)

                continue


            resposta.raise_for_status()


            # ==================================================
            # VERIFICAR PÁGINA DE MANUTENÇÃO
            # ==================================================

            texto_html = resposta.text.lower()

            if (
                "serviço temporariamente suspenso"
                in texto_html
                or
                "servico temporariamente suspenso"
                in texto_html
            ):

                print()
                print(
                    "OGol entrou em manutenção."
                )

                print(
                    "Salvando o progresso e encerrando..."
                )

                wb.save(arquivo_saida)
                wb.close()

                raise SystemExit


            # ==================================================
            # EXTRAIR DATA
            # ==================================================

            soup = BeautifulSoup(
                resposta.text,
                "html.parser"
            )

            texto = soup.get_text(
                " ",
                strip=True
            )

            datas = re.findall(
                r"\b\d{4}-\d{2}-\d{2}\b",
                texto
            )


            if datas:

                data_nascimento = datetime.strptime(
                    datas[0],
                    "%Y-%m-%d"
                )

                celula_data = ws.cell(
                    linha,
                    col_data
                )

                celula_data.value = data_nascimento
                celula_data.number_format = "DD/MM/YYYY"


                # Salva imediatamente
                wb.save(arquivo_saida)


                encontradas += 1
                sucesso = True


                print(
                    f"[{contador}/{total}] "
                    f"{nome} -> "
                    f"{data_nascimento.strftime('%d/%m/%Y')} ✓"
                )

                break


            else:

                nao_encontradas += 1
                sucesso = True

                print(
                    f"[{contador}/{total}] "
                    f"{nome} -> DATA NÃO ENCONTRADA"
                )

                break


        except requests.exceptions.RequestException as erro:

            print(
                f"[{contador}/{total}] "
                f"{nome} -> erro de conexão "
                f"(tentativa {tentativa})"
            )

            if tentativa < tentativas_maximas:
                time.sleep(10)


    if not sucesso:

        erros += 1

        print(
            f"[{contador}/{total}] "
            f"{nome} -> ERRO APÓS "
            f"{tentativas_maximas} TENTATIVAS"
        )


    time.sleep(pausa_entre_jogadores)


# ==========================================================
# SALVAMENTO FINAL
# ==========================================================

wb.save(arquivo_saida)
wb.close()


# ==========================================================
# RESUMO
# ==========================================================

print()
print("=" * 60)
print("COLETA FINALIZADA")
print("=" * 60)

print(f"Pendentes processados: {total}")
print(f"Novas datas encontradas: {encontradas}")
print(f"Datas não encontradas: {nao_encontradas}")
print(f"Erros: {erros}")

print()
print("Arquivo:")
print(arquivo_saida)

print("=" * 60)