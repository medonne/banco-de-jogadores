import re
import time
import random
import unicodedata
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pathlib import Path

# =========================================================
# CONFIGURAÇÕES
# =========================================================

arquivo_entrada = "Banco_Jogadores_FINAL.xlsx"
arquivo_saida = "Banco_Jogadores_ESTADOS.xlsx"

TEMPO_MINIMO = 2.0
TEMPO_MAXIMO = 4.0

ESTADOS_BRASIL = {
    "AC", "AL", "AP", "AM", "BA", "CE", "DF",
    "ES", "GO", "MA", "MT", "MS", "MG", "PA",
    "PB", "PR", "PE", "PI", "RJ", "RN", "RS",
    "RO", "RR", "SC", "SP", "SE", "TO"
}

# =========================================================
# FUNÇÕES AUXILIARES
# =========================================================

def normalizar(texto):
    """
    Remove acentos apenas para facilitar comparações.
    """
    texto = str(texto)

    texto = unicodedata.normalize("NFD", texto)

    texto = "".join(
        caractere
        for caractere in texto
        if unicodedata.category(caractere) != "Mn"
    )

    return texto.upper().strip()


def limpar_pais(pais):
    pais = re.sub(r"\s+", " ", pais).strip()

    return pais.upper()


# =========================================================
# LOCALIZA PLANILHA
# =========================================================

pasta_script = Path(__file__).resolve().parent

caminho_entrada = pasta_script / arquivo_entrada
caminho_saida = pasta_script / arquivo_saida

if not caminho_entrada.exists():
    print(f"ERRO: arquivo não encontrado:")
    print(caminho_entrada)

    input("\nPressione ENTER para sair...")
    raise SystemExit


# =========================================================
# ABRE EXCEL
# =========================================================

wb = load_workbook(caminho_entrada)
ws = wb.active

cabecalhos = {}

for celula in ws[1]:

    if celula.value:

        cabecalhos[
            str(celula.value).strip()
        ] = celula.column


col_nome = cabecalhos.get("Nome")
col_link = cabecalhos.get("Link_Ogol")
col_estado = cabecalhos.get("Estado_Origem")

if not col_nome:
    raise ValueError("Coluna Nome não encontrada.")

if not col_link:
    raise ValueError("Coluna Link_Ogol não encontrada.")

if not col_estado:
    raise ValueError("Coluna Estado_Origem não encontrada.")


# =========================================================
# CONEXÃO COM OGOL
# =========================================================

session = requests.Session()

session.headers.update({

    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/151.0.0.0 Safari/537.36"
    ),

    "Accept-Language":
        "pt-BR,pt;q=0.9,en-US;q=0.8,en;q=0.7"

})


# =========================================================
# BUSCAR ORIGEM
# =========================================================

def buscar_origem(url):

    try:

        resposta = session.get(
            url,
            timeout=20
        )

    except requests.RequestException as erro:

        print(f"   Erro de conexão: {erro}")

        return "ERRO_CONEXAO"


    if resposta.status_code in (
        403,
        429,
        500,
        502,
        503,
        504
    ):

        print(
            f"   OGol indisponível/bloqueando "
            f"(HTTP {resposta.status_code})"
        )

        return "PARAR"


    if resposta.status_code != 200:

        print(
            f"   HTTP {resposta.status_code}"
        )

        return None


    soup = BeautifulSoup(
        resposta.text,
        "html.parser"
    )


    # Texto completo da página
    texto = soup.get_text(
        " ",
        strip=True
    )

    texto_normalizado = normalizar(texto)


    # =====================================================
    # LOCALIZA SOMENTE A ÁREA DA NATURALIDADE
    #
    # Exemplo:
    #
    # PAÍS DE NASCIMENTO (NATURALIDADE)
    # Brasil (Lagoa Vermelha (RS))
    # POSIÇÃO
    # =====================================================

    inicio = texto_normalizado.find(
        "PAIS DE NASCIMENTO (NATURALIDADE)"
    )

    if inicio == -1:

        return None


    trecho_original = texto[
        inicio:
        inicio + 350
    ]

    trecho_normalizado = normalizar(
        trecho_original
    )


    # =====================================================
    # PRIMEIRO TENTA IDENTIFICAR UF
    # =====================================================

    ufs_encontradas = re.findall(
        r"\(([A-Z]{2})\)",
        trecho_normalizado
    )

    for uf in reversed(ufs_encontradas):

        if uf in ESTADOS_BRASIL:

            return uf


    # =====================================================
    # IDENTIFICA O PAÍS
    # =====================================================

    depois_label = re.split(
        r"PAIS DE NASCIMENTO\s*\(NATURALIDADE\)",
        trecho_normalizado,
        maxsplit=1
    )

    if len(depois_label) < 2:

        return None


    conteudo = depois_label[1].strip()


    # Para antes dos próximos campos da página
    conteudo = re.split(
        r"\bPOSICAO\b"
        r"|\bPE PREFERENCIAL\b"
        r"|\bALTURA\b"
        r"|\bSITUACAO\b",
        conteudo
    )[0].strip()


    # =====================================================
    # DESCOBRE O PAÍS
    #
    # Exemplos:
    #
    # BRASIL (LAGOA VERMELHA (RS))
    # ARGENTINA (BUENOS AIRES)
    # URUGUAI (MONTEVIDEO)
    # PORTUGAL (LISBOA)
    # =====================================================

    pais_match = re.match(
        r"^([A-ZÀ-Ü\s\-]+?)(?=\s*\(|$)",
        conteudo
    )


    if not pais_match:

        return None


    pais = limpar_pais(
        pais_match.group(1)
    )


    # =====================================================
    # REGRA PRINCIPAL
    # =====================================================

    if pais == "BRASIL":

        # Brasileiro sem UF encontrada:
        # NÃO grava "BRASIL".
        return None


    # Estrangeiro:
    # grava o país
    return pais


# =========================================================
# PROCESSAMENTO
# =========================================================

total = 0
mantidos = 0
revisados = 0
corrigidos = 0
nao_encontrados = 0
sem_link = 0


for linha in range(
    2,
    ws.max_row + 1
):

    nome = ws.cell(
        linha,
        col_nome
    ).value

    link = ws.cell(
        linha,
        col_link
    ).value

    origem_atual = ws.cell(
        linha,
        col_estado
    ).value


    if not nome:
        continue


    total += 1


    origem_texto = ""

    if origem_atual is not None:

        origem_texto = str(
            origem_atual
        ).strip().upper()


    # =====================================================
    # PRESERVA O QUE JÁ ESTÁ CORRETO
    #
    # Exemplo:
    # RS
    # SP
    # MG
    # ARGENTINA
    # URUGUAI
    #
    # Só revisa:
    # vazio ou BRASIL
    # =====================================================

    if (
        origem_texto
        and origem_texto != "BRASIL"
    ):

        mantidos += 1
        continue


    revisados += 1


    if not link:

        print(
            f"[{total}] {nome}: SEM LINK OGOL"
        )

        sem_link += 1

        continue


    print(
        f"[{total}] Revisando: {nome}"
    )


    resultado = buscar_origem(
        str(link).strip()
    )


    # =====================================================
    # SITE BLOQUEOU
    # =====================================================

    if resultado == "PARAR":

        print(
            "\nOGol ficou indisponível."
        )

        print(
            "Salvando o progresso..."
        )

        wb.save(
            caminho_saida
        )

        break


    # =====================================================
    # ERRO DE CONEXÃO
    # =====================================================

    elif resultado == "ERRO_CONEXAO":

        print(
            "   Erro de conexão."
        )


    # =====================================================
    # ENCONTROU
    # =====================================================

    elif resultado:

        ws.cell(
            linha,
            col_estado
        ).value = resultado

        corrigidos += 1

        print(
            f"   Origem correta: {resultado}"
        )

        # Salva a cada jogador encontrado
        wb.save(
            caminho_saida
        )


    # =====================================================
    # NÃO ENCONTROU
    # =====================================================

    else:

        # Se estava escrito BRASIL incorretamente,
        # apaga e deixa vazio.
        if origem_texto == "BRASIL":

            ws.cell(
                linha,
                col_estado
            ).value = None

        nao_encontrados += 1

        print(
            "   Estado não identificado."
        )

        wb.save(
            caminho_saida
        )


    time.sleep(
        random.uniform(
            TEMPO_MINIMO,
            TEMPO_MAXIMO
        )
    )


else:

    wb.save(
        caminho_saida
    )


# =========================================================
# RESUMO
# =========================================================

print("\n================================")
print("PROCESSO FINALIZADO")
print("================================")

print(
    f"Jogadores na planilha: {total}"
)

print(
    f"Origens já corretas e mantidas: {mantidos}"
)

print(
    f"Registros revisados: {revisados}"
)

print(
    f"Origens corrigidas/encontradas: {corrigidos}"
)

print(
    f"Origem não encontrada: {nao_encontrados}"
)

print(
    f"Sem link OGol: {sem_link}"
)

print("\nArquivo criado:")

print(
    caminho_saida
)