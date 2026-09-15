import re
import time
import random
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pathlib import Path

# =========================================================
# CONFIGURAÇÕES
# =========================================================

arquivo = "Banco_Jogadores_Cristiano.xlsx"

TEMPO_MINIMO = 2.0
TEMPO_MAXIMO = 4.0

ESTADOS_BRASIL = {
    "AC", "AL", "AP", "AM", "BA", "CE", "DF",
    "ES", "GO", "MA", "MT", "MS", "MG", "PA",
    "PB", "PR", "PE", "PI", "RJ", "RN", "RS",
    "RO", "RR", "SC", "SP", "SE", "TO"
}

# =========================================================
# LOCALIZA O ARQUIVO
# =========================================================

pasta = Path(__file__).resolve().parent
caminho = pasta / arquivo

if not caminho.exists():
    print(f"ERRO: arquivo não encontrado:")
    print(caminho)
    input("\nPressione ENTER para sair...")
    raise SystemExit

# =========================================================
# ABRE O EXCEL
# =========================================================

wb = load_workbook(caminho)
ws = wb.active

cabecalhos = {}

for celula in ws[1]:
    if celula.value:
        cabecalhos[str(celula.value).strip()] = celula.column

col_nome = cabecalhos.get("Nome")
col_link = cabecalhos.get("Link_Ogol")
col_estado = cabecalhos.get("Estado_Origem")

if not col_nome:
    raise ValueError("Coluna 'Nome' não encontrada.")

if not col_link:
    raise ValueError("Coluna 'Link_Ogol' não encontrada.")

if not col_estado:
    raise ValueError("Coluna 'Estado_Origem' não encontrada.")

# =========================================================
# CONEXÃO
# =========================================================

session = requests.Session()

session.headers.update({
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/151.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "pt-BR,pt;q=0.9"
})

# =========================================================
# FUNÇÃO PARA BUSCAR A ORIGEM
# =========================================================

def buscar_origem(url):

    try:
        resposta = session.get(url, timeout=20)

    except requests.RequestException as erro:
        print(f"   Erro de conexão: {erro}")
        return "ERRO_CONEXAO"

    if resposta.status_code in (403, 429, 500, 502, 503, 504):
        print(f"   OGol indisponível/bloqueando. HTTP {resposta.status_code}")
        return "PARAR"

    if resposta.status_code != 200:
        print(f"   HTTP {resposta.status_code}")
        return None

    soup = BeautifulSoup(
        resposta.text,
        "html.parser"
    )

    texto = soup.get_text(
        "\n",
        strip=True
    )

    linhas = [
        linha.strip()
        for linha in texto.splitlines()
        if linha.strip()
    ]

    # Procura "País de Nascimento (Naturalidade)"
    for i, linha in enumerate(linhas):

        if "PAÍS DE NASCIMENTO" in linha.upper():

            if i + 1 >= len(linhas):
                return None

            naturalidade = linhas[i + 1].strip()

            print(f"   OGol: {naturalidade}")

            # =================================================
            # BRASILEIRO
            # =================================================

            if naturalidade.upper().startswith("BRASIL"):

                # Primeiro procura UF.
                # Exemplo:
                # Brasil (Lagoa Vermelha (RS))
                match_uf = re.search(
                    r"\(([A-Z]{2})\)\s*\)?$",
                    naturalidade,
                    flags=re.IGNORECASE
                )

                if match_uf:

                    uf = match_uf.group(1).upper()

                    if uf in ESTADOS_BRASIL:
                        return uf

                # =============================================
                # NÃO TEM UF -> RETORNA A CIDADE
                #
                # Brasil (Suzano)
                # vira:
                # Suzano
                # =============================================

                match_cidade = re.match(
                    r"^Brasil\s*\((.+)\)\s*$",
                    naturalidade,
                    flags=re.IGNORECASE
                )

                if match_cidade:

                    cidade = match_cidade.group(1).strip()

                    return cidade

                return None

            # =================================================
            # ESTRANGEIRO
            # =================================================

            match_pais = re.match(
                r"^([^(]+)",
                naturalidade
            )

            if match_pais:

                pais = match_pais.group(1).strip().upper()

                return pais

            return None

    return None

# =========================================================
# PROCESSAMENTO
# =========================================================

total = 0
atualizados = 0
mantidos = 0
nao_encontrados = 0

for linha in range(2, ws.max_row + 1):

    nome = ws.cell(linha, col_nome).value
    link = ws.cell(linha, col_link).value
    origem_atual = ws.cell(linha, col_estado).value

    if not nome:
        continue

    total += 1

    origem_texto = ""

    if origem_atual is not None:
        origem_texto = str(origem_atual).strip()

    # =====================================================
    # IMPORTANTE:
    # Só busca quem estiver:
    #
    # vazio
    # ou
    # BRASIL
    #
    # Não mexe em RS, SP, MG, ARGENTINA etc.
    # =====================================================

    if origem_texto and origem_texto.upper() != "BRASIL":
        mantidos += 1
        continue

    if not link:
        print(f"[{total}] {nome}: sem Link_Ogol")
        nao_encontrados += 1
        continue

    print(f"\n[{total}] {nome}")

    resultado = buscar_origem(
        str(link).strip()
    )

    if resultado == "PARAR":

        print("\nOGol ficou indisponível.")
        print("Salvando o progresso...")

        wb.save(caminho)

        break

    elif resultado == "ERRO_CONEXAO":

        print("   Não foi possível acessar.")

        nao_encontrados += 1

    elif resultado:

        ws.cell(
            linha,
            col_estado
        ).value = resultado

        atualizados += 1

        print(f"   SALVO NO EXCEL: {resultado}")

        # Salva imediatamente
        wb.save(caminho)

    else:

        print("   Origem não encontrada.")

        nao_encontrados += 1

    time.sleep(
        random.uniform(
            TEMPO_MINIMO,
            TEMPO_MAXIMO
        )
    )

else:
    wb.save(caminho)

# =========================================================
# RESUMO
# =========================================================

print("\n================================")
print("PROCESSO FINALIZADO")
print("================================")

print(f"Jogadores: {total}")
print(f"Já estavam preenchidos: {mantidos}")
print(f"Atualizados: {atualizados}")
print(f"Não encontrados: {nao_encontrados}")

print("\nArquivo atualizado:")
print(caminho)