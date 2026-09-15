import time
import random
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pathlib import Path


# =========================================================
# CONFIGURAÇÕES
# =========================================================

arquivo_entrada = "Banco_Jogadores_CLUBES_ATUALIZADOS.xlsx"
arquivo_saida = "Banco_Jogadores_CLUBES_FINAL.xlsx"

INICIAR_NO_JOGADOR = 306

TEMPO_MINIMO = 2.0
TEMPO_MAXIMO = 4.0


# =========================================================
# CAMINHOS
# =========================================================

pasta = Path(__file__).resolve().parent

caminho_entrada = pasta / arquivo_entrada
caminho_saida = pasta / arquivo_saida


if not caminho_entrada.exists():

    print("Arquivo de entrada não encontrado:")
    print(caminho_entrada)

    raise SystemExit


# =========================================================
# ABRE PLANILHA
# =========================================================

wb = load_workbook(caminho_entrada)
ws = wb.active


# =========================================================
# LOCALIZA COLUNAS
# =========================================================

cabecalhos = {}

for celula in ws[1]:

    if celula.value:

        cabecalhos[
            str(celula.value).strip()
        ] = celula.column


col_nome = cabecalhos.get("Nome")
col_clube = cabecalhos.get("Clube_Atual")
col_pais = cabecalhos.get("Pais")
col_link = cabecalhos.get("Link_Ogol")


# =========================================================
# CONEXÃO OGOL
# =========================================================

session = requests.Session()

session.headers.update({

    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/151.0.0.0 Safari/537.36"
    ),

    "Accept-Language": "pt-BR,pt;q=0.9"

})


# =========================================================
# FUNÇÃO
# =========================================================

def buscar_clube_pais(url):

    try:

        resposta = session.get(
            url,
            timeout=20
        )

    except requests.RequestException as erro:

        print(
            f"   Erro de conexão: {erro}"
        )

        return None, None


    if resposta.status_code != 200:

        print(
            f"   HTTP {resposta.status_code}"
        )

        return None, None


    soup = BeautifulSoup(
        resposta.text,
        "html.parser"
    )


    linhas = soup.select(
        ".card-data__row"
    )


    for linha in linhas:

        label = linha.select_one(
            ".card-data__label"
        )

        if not label:
            continue


        nome_label = label.get_text(
            " ",
            strip=True
        ).lower()


        if nome_label != "clube atual":
            continue


        # =================================================
        # CLUBE
        # =================================================

        clube = None

        texto_clube = linha.select_one(
            ".micrologo_and_text .text"
        )

        if texto_clube:

            clube = texto_clube.get_text(
                " ",
                strip=True
            )


        # =================================================
        # PAÍS
        # =================================================

        pais = None

        link_pais = linha.select_one(
            'a[href*="/pais/"]'
        )


        if link_pais:

            pais = link_pais.get(
                "title"
            )


            if not pais:

                href = link_pais.get(
                    "href",
                    ""
                )

                if "/pais/" in href:

                    pais = href.split(
                        "/pais/"
                    )[-1]

                    pais = pais.replace(
                        "-",
                        " "
                    ).title()


        return clube, pais


    return None, None


# =========================================================
# PROCESSAMENTO
# =========================================================

total = 0
atualizados = 0
nao_encontrados = 0


for linha in range(
    2,
    ws.max_row + 1
):

    nome = ws.cell(
        linha,
        col_nome
    ).value


    if not nome:
        continue


    total += 1


    # =====================================================
    # PULA OS 305 QUE JÁ FORAM FEITOS
    # =====================================================

    if total < INICIAR_NO_JOGADOR:
        continue


    link = ws.cell(
        linha,
        col_link
    ).value


    print("\n--------------------------------")

    print(
        f"[{total}] {nome}"
    )


    if not link:

        print(
            "   Sem Link_Ogol"
        )

        nao_encontrados += 1

        continue


    clube_antigo = ws.cell(
        linha,
        col_clube
    ).value

    pais_antigo = ws.cell(
        linha,
        col_pais
    ).value


    print(
        f"   Clube antigo: {clube_antigo}"
    )

    print(
        f"   País antigo: {pais_antigo}"
    )


    clube_novo, pais_novo = buscar_clube_pais(
        str(link).strip()
    )


    if clube_novo:

        clube_novo = clube_novo.strip()

        ws.cell(
            linha,
            col_clube
        ).value = clube_novo


        print(
            f"   Clube atual OGol: {clube_novo}"
        )


        if pais_novo:

            pais_novo = pais_novo.strip().upper()

            ws.cell(
                linha,
                col_pais
            ).value = pais_novo


            print(
                f"   País do clube: {pais_novo}"
            )


        atualizados += 1


    else:

        print(
            "   Clube atual não encontrado."
        )

        nao_encontrados += 1


    # =====================================================
    # SALVA NO NOVO ARQUIVO
    # =====================================================

    try:

        wb.save(
            caminho_saida
        )

    except PermissionError:

        print(
            "\nERRO: o arquivo de saída ficou bloqueado."
        )

        print(
            "O progresso permanece na memória, mas o Python será encerrado."
        )

        raise


    time.sleep(
        random.uniform(
            TEMPO_MINIMO,
            TEMPO_MAXIMO
        )
    )


# =========================================================
# SALVA FINAL
# =========================================================

wb.save(
    caminho_saida
)


print("\n================================")
print("PROCESSO FINALIZADO")
print("================================")

print(
    f"Último jogador: {total}"
)

print(
    f"Atualizados nesta continuação: {atualizados}"
)

print(
    f"Não encontrados: {nao_encontrados}"
)

print("\nArquivo final:")

print(
    caminho_saida
)