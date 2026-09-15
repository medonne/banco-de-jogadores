import requests
import time
import random

from bs4 import BeautifulSoup
from openpyxl import load_workbook, Workbook
from pathlib import Path
from urllib.parse import urljoin


# =========================================================
# CONFIGURAÇÕES
# =========================================================

arquivo_entrada = "Banco_Jogadores_FINAL_ATUALIZADO.xlsx"
arquivo_saida = "Revisar_Campeonatos.xlsx"

TEMPO_MINIMO = 2.0
TEMPO_MAXIMO = 4.0


# =========================================================
# CAMINHOS
# =========================================================

pasta = Path(__file__).resolve().parent

caminho_entrada = pasta / arquivo_entrada
caminho_saida = pasta / arquivo_saida


# =========================================================
# ABRE PLANILHA
# =========================================================

wb = load_workbook(caminho_entrada)
ws = wb.active


cabecalhos = {}

for celula in ws[1]:
    if celula.value:
        cabecalhos[str(celula.value).strip()] = celula.column


col_id = cabecalhos.get("ID_Jogador")
col_nome = cabecalhos.get("Nome")
col_clube = cabecalhos.get("Clube_Atual")
col_campeonato = cabecalhos.get("Campeonato")
col_pais = cabecalhos.get("Pais")
col_link = cabecalhos.get("Link_Ogol")


# =========================================================
# CONEXÃO
# =========================================================

session = requests.Session()

session.headers.update({
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/151.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "pt-BR,pt;q=0.9"
})


# =========================================================
# PROCURA EDIÇÃO DO CLUBE ATUAL
# =========================================================

def buscar_edicao(url_jogador, clube_atual):

    try:
        resposta = session.get(
            url_jogador,
            timeout=20
        )

    except requests.RequestException:
        return None

    if resposta.status_code != 200:
        return None

    soup = BeautifulSoup(
        resposta.text,
        "html.parser"
    )

    clube_lower = clube_atual.lower()

    candidatos = []

    for link in soup.find_all("a", href=True):

        href = link.get("href", "")

        if "edition.php?id_edicao=" not in href:
            continue

        elemento = link
        encontrou_clube = False

        for _ in range(8):

            elemento = elemento.parent

            if elemento is None:
                break

            texto_bloco = elemento.get_text(
                " ",
                strip=True
            )

            if clube_lower in texto_bloco.lower():
                encontrou_clube = True
                break

        if encontrou_clube:

            url_edicao = urljoin(
                url_jogador,
                href
            )

            if url_edicao not in candidatos:
                candidatos.append(url_edicao)

    if candidatos:
        return candidatos[0]

    return None


# =========================================================
# CRIA PLANILHA DE REVISÃO
# =========================================================

wb_saida = Workbook()
ws_saida = wb_saida.active

ws_saida.title = "Revisar Campeonatos"

ws_saida.append([
    "ID_Jogador",
    "Nome",
    "Clube_Atual",
    "Pais",
    "Campeonato_Atual_Planilha",
    "Link_Ogol"
])


# =========================================================
# PROCESSAMENTO
# =========================================================

total = 0
nao_encontrados = 0


for linha in range(2, ws.max_row + 1):

    nome = ws.cell(
        linha,
        col_nome
    ).value

    clube = ws.cell(
        linha,
        col_clube
    ).value

    link = ws.cell(
        linha,
        col_link
    ).value

    if not nome:
        continue

    total += 1

    print("\n================================")
    print(f"[{total}] {nome}")
    print(f"Clube: {clube}")

    if not clube or not link:
        continue

    edicao = buscar_edicao(
        str(link).strip(),
        str(clube).strip()
    )

    if not edicao:

        nao_encontrados += 1

        print(">>> CAMPEONATO NÃO ENCONTRADO")

        ws_saida.append([
            ws.cell(linha, col_id).value,
            nome,
            clube,
            ws.cell(linha, col_pais).value,
            ws.cell(linha, col_campeonato).value,
            link
        ])

        wb_saida.save(
            caminho_saida
        )

    else:
        print("OK")

    time.sleep(
        random.uniform(
            TEMPO_MINIMO,
            TEMPO_MAXIMO
        )
    )


# =========================================================
# FINAL
# =========================================================

wb_saida.save(
    caminho_saida
)

print("\n================================")
print("PROCESSO FINALIZADO")
print("================================")

print(f"Jogadores analisados: {total}")
print(f"Não encontrados: {nao_encontrados}")

print("\nArquivo criado:")
print(caminho_saida)