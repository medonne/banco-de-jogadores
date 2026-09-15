import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pathlib import Path


# =========================================================
# ARQUIVO
# =========================================================

pasta = Path(__file__).resolve().parent
arquivo = pasta / "Revisar_Campeonatos.xlsx"

wb = load_workbook(arquivo)
ws = wb.active


# =========================================================
# CABEÇALHOS
# =========================================================

cabecalhos = {}

for celula in ws[1]:
    if celula.value:
        cabecalhos[str(celula.value).strip()] = celula.column


col_nome = cabecalhos["Nome"]
col_clube = cabecalhos["Clube_Atual"]
col_link = cabecalhos["Link_Ogol"]


# =========================================================
# PRIMEIRO JOGADOR DA LISTA
# =========================================================

linha = 2

nome = ws.cell(linha, col_nome).value
clube = ws.cell(linha, col_clube).value
url_jogador = ws.cell(linha, col_link).value


print("================================")
print("JOGADOR")
print("================================")

print("Nome:", nome)
print("Clube atual:", clube)
print("Link:", url_jogador)


# =========================================================
# ACESSA OGOL
# =========================================================

headers = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/151.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "pt-BR,pt;q=0.9"
}


resposta = requests.get(
    url_jogador,
    headers=headers,
    timeout=20
)

print("\nHTTP:", resposta.status_code)


soup = BeautifulSoup(
    resposta.text,
    "html.parser"
)


# =========================================================
# PROCURA LINKS COM O NOME DO CLUBE
# =========================================================

print("\n================================")
print("LINKS ENCONTRADOS PARA O CLUBE")
print("================================")


encontrados = 0


for link in soup.find_all("a", href=True):

    texto = link.get_text(
        " ",
        strip=True
    )

    if not texto:
        continue


    if clube.lower() in texto.lower():

        encontrados += 1

        print("\n------------------------------")
        print("Texto:", texto)
        print("Href:", link["href"])


print("\n================================")
print("TOTAL:", encontrados)
print("================================")