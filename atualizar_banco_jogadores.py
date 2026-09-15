import re
import time
import random
import requests

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pathlib import Path
from urllib.parse import urljoin


# =========================================================
# CONFIGURAÇÕES
# =========================================================

ARQUIVO_ENTRADA = "Banco_Jogadores.xlsx"
ARQUIVO_SAIDA = "Banco_Jogadores_Atualizado.xlsx"

TEMPO_MINIMO = 2.0
TEMPO_MAXIMO = 4.0

ANO_ATUAL = "2026"

# Na estrutura atual do OGol:
# epoca_id=155 corresponde à temporada 2026.
EPOCA_ID_ATUAL = "155"


# =========================================================
# CAMINHOS
# =========================================================

pasta = Path(__file__).resolve().parent

caminho_entrada = pasta / ARQUIVO_ENTRADA
caminho_saida = pasta / ARQUIVO_SAIDA


if not caminho_entrada.exists():
    print("ERRO: arquivo de entrada não encontrado:")
    print(caminho_entrada)
    raise SystemExit


# =========================================================
# EXCEL
# =========================================================

wb = load_workbook(caminho_entrada)
ws = wb.active


cabecalhos = {}

for celula in ws[1]:

    if celula.value:
        cabecalhos[str(celula.value).strip()] = celula.column


col_id = cabecalhos.get("ID_Jogador")
col_nome = cabecalhos.get("Nome")
col_clube = cabecalhos.get("Clube_Atual")
col_campeonato = cabecalhos.get("Campeonato")
col_pais = cabecalhos.get("Pais")
col_link = cabecalhos.get("Link_Ogol")


obrigatorias = {
    "ID_Jogador": col_id,
    "Nome": col_nome,
    "Clube_Atual": col_clube,
    "Campeonato": col_campeonato,
    "Pais": col_pais,
    "Link_Ogol": col_link
}


faltando = [
    nome
    for nome, coluna in obrigatorias.items()
    if coluna is None
]


if faltando:

    print("ERRO: colunas não encontradas:")
    print(", ".join(faltando))

    raise SystemExit


# =========================================================
# CONEXÃO
# =========================================================

session = requests.Session()

session.headers.update({

    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/151.0.0.0 Safari/537.36"
    ),

    "Accept-Language": "pt-BR,pt;q=0.9"
})


# =========================================================
# FUNÇÃO DE REQUISIÇÃO
# =========================================================

def baixar_pagina(url):

    try:

        resposta = session.get(
            url,
            timeout=20
        )

    except requests.RequestException as erro:

        print("   Erro de conexão:", erro)
        return None


    if resposta.status_code != 200:

        print(
            f"   HTTP: {resposta.status_code}"
        )

        return None


    return BeautifulSoup(
        resposta.text,
        "html.parser"
    )


# =========================================================
# CLUBE ATUAL + PAÍS
# =========================================================

def buscar_clube_e_pais(url_jogador):

    soup = baixar_pagina(
        url_jogador
    )


    if soup is None:

        return None, None, None


    clube_atual = None
    pais = None


    for linha in soup.select(
        ".card-data__row"
    ):

        label = linha.select_one(
            ".card-data__label"
        )


        if not label:
            continue


        nome_label = label.get_text(
            " ",
            strip=True
        ).lower()


        if nome_label != "clube atual":
            continue


        clube = linha.select_one(
            ".micrologo_and_text .text"
        )


        if clube:

            clube_atual = clube.get_text(
                " ",
                strip=True
            )


        link_pais = linha.select_one(
            'a[href*="/pais/"]'
        )


        if link_pais:

            pais = link_pais.get(
                "title"
            )


            if not pais:

                href = link_pais.get(
                    "href",
                    ""
                )


                if "/pais/" in href:

                    pais = href.split(
                        "/pais/"
                    )[-1].split("?")[0]


        break


    if pais:
        pais = pais.strip().upper()


    return clube_atual, pais, soup


# =========================================================
# PADRONIZA NOME DO CAMPEONATO
# =========================================================

def limpar_campeonato(nome):

    if not nome:
        return None


    nome = nome.strip()


    # remove ano final
    nome = re.sub(
        r"\s+20\d{2}\s*$",
        "",
        nome
    )


    nome_upper = nome.upper()


    # =====================================================
    # BRASILEIRO
    # =====================================================

    if (
        "BRASILEIRÃO" in nome_upper
        and "SÉRIE A" in nome_upper
    ):
        return "SÉRIE A"


    if (
        "BRASILEIRÃO" in nome_upper
        and "SÉRIE B" in nome_upper
    ):
        return "SÉRIE B"


    if (
        "BRASILEIRÃO" in nome_upper
        and "SÉRIE C" in nome_upper
    ):
        return "SÉRIE C"


    if (
        "BRASILEIRÃO" in nome_upper
        and "SÉRIE D" in nome_upper
    ):
        return "SÉRIE D"


    # =====================================================
    # GAÚCHO ACESSO
    # =====================================================

    if (
        "GAÚCHO" in nome_upper
        and (
            "DIVISÃO DE ACESSO" in nome_upper
            or "ACESSO" in nome_upper
        )
    ):
        return "GAÚCHO ACESSO"


    # remove "Campeonato" do início

    nome = re.sub(
        r"^Campeonato\s+",
        "",
        nome,
        flags=re.IGNORECASE
    )


    return nome.upper()


# =========================================================
# NOME DA PÁGINA DE UMA EDIÇÃO
# =========================================================

def buscar_nome_edicao(url_edicao):

    soup = baixar_pagina(
        url_edicao
    )


    if soup is None:
        return None


    h1 = soup.find("h1")


    if h1:

        nome = h1.get_text(
            " ",
            strip=True
        )


        if nome:

            return limpar_campeonato(
                nome
            )


    if soup.title:

        titulo = soup.title.get_text(
            " ",
            strip=True
        )

        titulo = titulo.split(
            " - "
        )[0]


        return limpar_campeonato(
            titulo
        )


    return None


# =========================================================
# MÉTODO 1:
# DESCOBRE CAMPEONATO PELOS JOGOS DO CLUBE ATUAL
# =========================================================

def campeonato_pelos_jogos(
    soup_jogador,
    url_jogador,
    clube_atual
):

    if soup_jogador is None:
        return None


    clube_lower = clube_atual.lower()

    candidatos = []


    for link in soup_jogador.find_all(
        "a",
        href=True
    ):

        href = link.get(
            "href",
            ""
        )


        if "edition.php?id_edicao=" not in href:
            continue


        elemento = link
        encontrou_clube = False


        for _ in range(8):

            elemento = elemento.parent


            if elemento is None:
                break


            texto_bloco = elemento.get_text(
                " ",
                strip=True
            )


            if clube_lower in texto_bloco.lower():

                encontrou_clube = True
                break


        if encontrou_clube:

            url_edicao = urljoin(
                url_jogador,
                href
            )


            if url_edicao not in candidatos:

                candidatos.append(
                    url_edicao
                )


    for url_edicao in candidatos:

        campeonato = buscar_nome_edicao(
            url_edicao
        )


        if campeonato:

            return campeonato


    return None


# =========================================================
# DESCOBRE LINK DO CLUBE
# =========================================================

def buscar_link_clube(
    soup_jogador,
    url_jogador,
    clube_atual
):

    candidatos = []


    for link in soup_jogador.find_all(
        "a",
        href=True
    ):

        texto = link.get_text(
            " ",
            strip=True
        )


        href = link.get(
            "href",
            ""
        )


        if not texto:
            continue


        if clube_atual.lower() not in texto.lower():
            continue


        if "/equipe/" not in href:
            continue


        url = urljoin(
            url_jogador,
            href
        )


        if url not in candidatos:

            candidatos.append(
                url
            )


    if not candidatos:

        return None


    # tenta temporada atual primeiro

    for url in candidatos:

        if (
            f"epoca_id={EPOCA_ID_ATUAL}"
            in url
        ):

            return url


    # depois tenta link sem temporada

    for url in candidatos:

        if "epoca_id=" not in url:

            return url


    return candidatos[0]


# =========================================================
# PEGA TODAS AS COMPETIÇÕES DO CLUBE
# =========================================================

def buscar_competicoes_clube(
    url_clube
):

    soup = baixar_pagina(
        url_clube
    )


    if soup is None:

        return []


    competicoes = []


    for link in soup.find_all(
        "a",
        href=True
    ):

        href = link.get(
            "href",
            ""
        )


        if "edition.php?id_edicao=" not in href:
            continue


        texto = link.get_text(
            " ",
            strip=True
        )


        if not texto:
            continue


        # só temporada atual
        if ANO_ATUAL not in texto:
            continue


        url_edicao = urljoin(
            url_clube,
            href
        )


        item = (
            texto,
            url_edicao
        )


        if item not in competicoes:

            competicoes.append(
                item
            )


    return competicoes


# =========================================================
# ESCOLHE CAMPEONATO PRINCIPAL DO CLUBE
# =========================================================

def escolher_campeonato_clube(
    competicoes,
    pais
):

    if not competicoes:

        return None


    pais_upper = str(
        pais or ""
    ).upper()


    # =====================================================
    # BRASIL:
    # prioridade para campeonato nacional
    # =====================================================

    if pais_upper == "BRASIL":

        prioridades = [
            "BRASILEIRÃO SÉRIE A",
            "BRASILEIRÃO SÉRIE B",
            "BRASILEIRÃO SÉRIE C",
            "BRASILEIRÃO SÉRIE D"
        ]


        for prioridade in prioridades:

            for nome, _ in competicoes:

                if prioridade in nome.upper():

                    return limpar_campeonato(
                        nome
                    )


    # =====================================================
    # EXCLUI COPAS E AMISTOSOS
    # =====================================================

    ignorar = [
        "COPA",
        "SUPERCOPA",
        "RECOPA",
        "AMISTOSO",
        "AMISTOSOS",
        "TAÇA"
    ]


    candidatos = []


    for nome, _ in competicoes:

        nome_upper = nome.upper()


        if any(
            palavra in nome_upper
            for palavra in ignorar
        ):

            continue


        candidatos.append(
            nome
        )


    if candidatos:

        return limpar_campeonato(
            candidatos[0]
        )


    return None


# =========================================================
# MÉTODO 2:
# CAMPEONATO PELA PÁGINA DO CLUBE
# =========================================================

def campeonato_pelo_clube(
    soup_jogador,
    url_jogador,
    clube_atual,
    pais
):

    link_clube = buscar_link_clube(
        soup_jogador,
        url_jogador,
        clube_atual
    )


    if not link_clube:

        return None


    competicoes = buscar_competicoes_clube(
        link_clube
    )


    return escolher_campeonato_clube(
        competicoes,
        pais
    )


# =========================================================
# PROCESSAMENTO
# =========================================================

total = 0
atualizados_clube = 0
atualizados_pais = 0
atualizados_campeonato = 0
campeonato_metodo_1 = 0
campeonato_metodo_2 = 0
nao_encontrados = 0


for linha in range(
    2,
    ws.max_row + 1
):

    nome = ws.cell(
        linha,
        col_nome
    ).value


    if not nome:
        continue


    total += 1


    link_jogador = ws.cell(
        linha,
        col_link
    ).value


    print("\n================================")
    print(f"[{total}] {nome}")
    print("================================")


    if not link_jogador:

        print("Sem Link_Ogol.")

        ws.cell(
            linha,
            col_campeonato
        ).value = "NÃO ENCONTRADO"

        nao_encontrados += 1

        continue


    # =====================================================
    # CLUBE + PAÍS
    # =====================================================

    clube_atual, pais, soup_jogador = buscar_clube_e_pais(
        str(link_jogador).strip()
    )


    if clube_atual:

        clube_antigo = ws.cell(
            linha,
            col_clube
        ).value


        ws.cell(
            linha,
            col_clube
        ).value = clube_atual


        atualizados_clube += 1


        print(
            f"Clube: {clube_antigo} -> {clube_atual}"
        )

    else:

        clube_atual = ws.cell(
            linha,
            col_clube
        ).value

        print(
            f"Clube não encontrado. Mantido: {clube_atual}"
        )


    if pais:

        pais_antigo = ws.cell(
            linha,
            col_pais
        ).value


        ws.cell(
            linha,
            col_pais
        ).value = pais


        atualizados_pais += 1


        print(
            f"País: {pais_antigo} -> {pais}"
        )

    else:

        pais = ws.cell(
            linha,
            col_pais
        ).value

        print(
            f"País não encontrado. Mantido: {pais}"
        )


    # =====================================================
    # CAMPEONATO
    # =====================================================

    campeonato = None


    if (
        clube_atual
        and soup_jogador is not None
    ):

        # -------------------------------------------------
        # MÉTODO 1
        # -------------------------------------------------

        campeonato = campeonato_pelos_jogos(
            soup_jogador,
            str(link_jogador).strip(),
            str(clube_atual).strip()
        )


        if campeonato:

            campeonato_metodo_1 += 1

            print(
                f"Campeonato pelo jogador: {campeonato}"
            )


        # -------------------------------------------------
        # MÉTODO 2
        # -------------------------------------------------

        if not campeonato:

            campeonato = campeonato_pelo_clube(
                soup_jogador,
                str(link_jogador).strip(),
                str(clube_atual).strip(),
                pais
            )


            if campeonato:

                campeonato_metodo_2 += 1

                print(
                    f"Campeonato pelo clube: {campeonato}"
                )


    # =====================================================
    # RESULTADO
    # =====================================================

    if campeonato:

        ws.cell(
            linha,
            col_campeonato
        ).value = campeonato


        atualizados_campeonato += 1


    else:

        ws.cell(
            linha,
            col_campeonato
        ).value = "NÃO ENCONTRADO"


        nao_encontrados += 1


        print(
            "Campeonato: NÃO ENCONTRADO"
        )


    # =====================================================
    # SALVA PROGRESSO
    # =====================================================

    try:

        wb.save(
            caminho_saida
        )

    except PermissionError:

        print("\nERRO: arquivo de saída bloqueado.")
        print("Feche o Excel e aguarde o OneDrive liberar o arquivo.")
        print("O programa será encerrado.")

        raise


    # =====================================================
    # INTERVALO
    # =====================================================

    time.sleep(
        random.uniform(
            TEMPO_MINIMO,
            TEMPO_MAXIMO
        )
    )


# =========================================================
# SALVA FINAL
# =========================================================

wb.save(
    caminho_saida
)


# =========================================================
# RESUMO
# =========================================================

print("\n================================")
print("PROCESSO FINALIZADO")
print("================================")

print(
    f"Jogadores analisados: {total}"
)

print(
    f"Clubes encontrados: {atualizados_clube}"
)

print(
    f"Países encontrados: {atualizados_pais}"
)

print(
    f"Campeonatos encontrados: {atualizados_campeonato}"
)

print(
    f"  Pelo jogador: {campeonato_metodo_1}"
)

print(
    f"  Pelo clube: {campeonato_metodo_2}"
)

print(
    f"Campeonatos NÃO ENCONTRADOS: {nao_encontrados}"
)

print("\nArquivo criado:")

print(
    caminho_saida
)